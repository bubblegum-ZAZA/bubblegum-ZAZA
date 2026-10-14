"""เติมแบบ ผส.2 (รายงานการฝึกหัดงานสอบบัญชี) และ ผส.3 (คำรับรองของผู้ให้การฝึกหัดงาน)

อ่านข้อมูลจากไฟล์ Excel แล้วพิมพ์ลงแบบฟอร์ม PDF ของสภาวิชาชีพบัญชี (ฉบับปรับปรุง 31 ต.ค. 2567)

    python fill_forms.py --make-template ข้อมูล.xlsx      # สร้างไฟล์ Excel เปล่าไว้กรอก
    python fill_forms.py ข้อมูล.xlsx                       # เติมฟอร์ม ได้ไฟล์ ผส.2 / ผส.3 ในโฟลเดอร์ output
"""

import argparse
import datetime as dt
import io
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pymupdf
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
BLANK_FORM = HERE / "forms" / "ps2-ps3-blank.pdf"

# หน้าในไฟล์ฟอร์มเปล่า
PAGE_PS3 = 0
PAGE_PS2 = 1
PAGE_DETAIL_SHORT = 2  # ชั่วโมงรวม <= 1,500
PAGE_DETAIL_LONG = 3  # ชั่วโมงรวม > 1,500
HOUR_LIMIT = 1500

# ใช้ฟอนต์ที่มากับ PyMuPDF (Noto Serif Thai + Nimbus Sans) ไม่ต้องติดตั้งฟอนต์เพิ่ม
# ฟอนต์ไทยบางตัว (เช่น Sarabun) ต้องยกวรรณยุกต์ด้วย GPOS ซึ่ง MuPDF วางไม่ถูก วรรณยุกต์บนสระจะหาย
FONT_SIZE = 11  # ≈ TH SarabunPSK 16 ที่ใช้ในฟอร์ม
TABLE_FONT_SIZE = 9.5
LINE_HEIGHT = 1.15
INK = "#10206b"  # น้ำเงินเข้มให้แยกออกจากตัวพิมพ์ของฟอร์ม

THAI_MONTHS = [
    "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]
THAI_MONTHS_SHORT = [
    "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
    "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.",
]

# ---------------------------------------------------------------- Excel

SHEET_INFO = "ข้อมูลทั่วไป"
SHEET_DETAIL = "รายละเอียดงาน"

# (key, คำอธิบาย, ตัวอย่าง)
INFO_FIELDS = [
    ("__ผู้ฝึกหัดงาน", None, None),
    ("เลขประจำตัวฝึกหัดงาน", "7 หลัก", "6712345"),
    ("ชื่อผู้ฝึกหัดงาน", "ใส่คำนำหน้าด้วย", "นางสาวสมหญิง ตั้งใจเรียน"),
    ("__ผู้ให้การฝึกหัดงาน (ผู้สอบบัญชีรับอนุญาต)", None, None),
    ("ชื่อผู้สอบบัญชี", "ใส่คำนำหน้าด้วย", "นายสมชาย ใจดี"),
    ("เลขทะเบียนผู้สอบบัญชี", "", "9999"),
    ("สำนักงาน", "", "สำนักงานสมชายการบัญชี"),
    ("เลขที่", "", "99/9"),
    ("หมู่ที่", "", "5"),
    ("ซอย", "", "-"),
    ("ถนน", "", "พหลโยธิน"),
    ("ตำบล/แขวง", "", "สามเสนใน"),
    ("อำเภอ/เขต", "", "พญาไท"),
    ("จังหวัด", "", "กรุงเทพมหานคร"),
    ("รหัสไปรษณีย์", "", "10400"),
    ("โทรศัพท์", "", "02-123-4567"),
    ("โทรสาร", "", "-"),
    ("E-mail", "", "somchai@example.com"),
    ("__ระยะเวลาฝึกหัดงาน", None, None),
    ("หักพักกลางวัน (ชม./วัน)", "ใช้ตอนคำนวณชั่วโมงจากเวลาเริ่ม-สิ้นสุด ที่คร่อมช่วง 12.00-13.00 น.", 1),
    ("ฝึกตั้งแต่วันที่", "เว้นว่าง = ใช้วันแรกในชีทรายละเอียดงาน", ""),
    ("ฝึกถึงวันที่", "เว้นว่าง = ใช้วันสุดท้ายในชีทรายละเอียดงาน", ""),
    ("__แบบ ผส.2", None, None),
    ("ผส.2 เขียนที่", "", "สำนักงานสมชายการบัญชี"),
    ("ผส.2 วันที่", "เว้นว่าง = วันนี้", ""),
    ("__แบบ ผส.3", None, None),
    ("ผส.3 เขียนที่", "เว้นว่าง = เหมือน ผส.2", ""),
    ("ผส.3 วันที่", "เว้นว่าง = เหมือน ผส.2", ""),
    ("ฝึกในฐานะ", "ข้อ 2 ของ ผส.3", "ผู้ช่วยผู้สอบบัญชี"),
    ("ความเห็นของผู้ให้การฝึก", "ต่อท้าย \"...เกี่ยวกับการสอบบัญชีและ\" ข้อ 3 ของ ผส.3",
     "สามารถปฏิบัติงานตรวจสอบบัญชีได้ด้วยตนเองตามมาตรฐานการสอบบัญชี"),
]

# (หัวคอลัมน์, ความกว้าง, ใช้กับแบบ)
DETAIL_COLUMNS = [
    ("ชื่อธุรกิจ", 32, "ทุกแบบ"),
    ("สถานที่ตั้ง", 36, "ทุกแบบ"),
    ("เบอร์โทรศัพท์", 14, "ทุกแบบ (ไม่มีเว้นว่าง)"),
    ("รอบปีบัญชี", 14, "ทุกแบบ เช่น 31/12/2566"),
    ("วันที่ผู้สอบบัญชีลงนามรับรองงบ", 16, "เฉพาะ > 1,500 ชม. (ยังไม่ลงนามเว้นว่าง)"),
    ("วันที่เริ่ม", 13, "วันที่ปฏิบัติงาน"),
    ("วันที่สิ้นสุด", 13, "เฉพาะ <= 1,500 ชม. (ทำวันเดียวเว้นว่าง)"),
    ("เวลาเริ่มต้น", 10, "เฉพาะ > 1,500 ชม. เช่น 08:30"),
    ("เวลาสิ้นสุด", 10, "เฉพาะ > 1,500 ชม. เช่น 17:30"),
    ("จำนวนชั่วโมง", 10, "เว้นว่าง = เวลาสิ้นสุด - เวลาเริ่มต้น - พักกลางวัน"),
    ("ลักษณะการปฏิบัติงาน", 30, "ตามคำชี้แจงข้อบังคับฯ (ฉบับที่ 3) พ.ศ. 2563"),
]

SAMPLE_ROWS = [
    ["บริษัท ตัวอย่าง จำกัด", "1 ถนนสุขุมวิท แขวงคลองเตย เขตคลองเตย กรุงเทพฯ 10110", "02-000-0001",
     dt.date(2023, 12, 31), dt.date(2024, 3, 15), dt.date(2024, 1, 8), None,
     dt.time(8, 30), dt.time(17, 30), 8, "ตรวจสอบเงินสดและเงินฝากธนาคาร"],
    ["บริษัท ตัวอย่าง จำกัด", "1 ถนนสุขุมวิท แขวงคลองเตย เขตคลองเตย กรุงเทพฯ 10110", "02-000-0001",
     dt.date(2023, 12, 31), dt.date(2024, 3, 15), dt.date(2024, 1, 9), None,
     dt.time(8, 30), dt.time(17, 30), 8, "ตรวจสอบลูกหนี้การค้า"],
    ["ห้างหุ้นส่วนจำกัด ทดลอง", "2 หมู่ 3 ตำบลบางพูด อำเภอปากเกร็ด นนทบุรี 11120", None,
     dt.date(2023, 12, 31), None, dt.date(2024, 1, 15), None,
     dt.time(9, 0), dt.time(16, 0), None, "ตรวจนับสินค้าคงเหลือ"],
]


def make_template(path: Path) -> None:
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="DDE7F7")
    sect_fill = PatternFill("solid", fgColor="F3E2B3")
    bold = Font(bold=True)

    ws = wb.active
    ws.title = SHEET_INFO
    ws.append(["หัวข้อ", "ข้อมูล", "หมายเหตุ"])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for key, note, sample in INFO_FIELDS:
        if key.startswith("__"):
            ws.append([key[2:]])
            ws.cell(ws.max_row, 1).font = bold
            for col in range(1, 4):
                ws.cell(ws.max_row, col).fill = sect_fill
        else:
            ws.append([key, sample, note])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 55
    for row in ws.iter_rows(min_col=2, max_col=2):
        row[0].number_format = "@"  # เลขประจำตัว / รหัสไปรษณีย์ ไม่ให้ Excel ตัดเลข 0 หน้า

    wd = wb.create_sheet(SHEET_DETAIL)
    wd.append([c[0] for c in DETAIL_COLUMNS])
    wd.append([c[2] for c in DETAIL_COLUMNS])
    for i, (_, width, _) in enumerate(DETAIL_COLUMNS, start=1):
        letter = openpyxl.utils.get_column_letter(i)
        wd.column_dimensions[letter].width = width
        wd.cell(1, i).font, wd.cell(1, i).fill = bold, head_fill
        wd.cell(2, i).font = Font(italic=True, color="808080", size=9)
        wd.cell(2, i).alignment = Alignment(wrap_text=True, vertical="top")
    for r in SAMPLE_ROWS:
        wd.append(r)
    for row in wd.iter_rows(min_row=3, max_row=500):
        for i, cell in enumerate(row):
            name = DETAIL_COLUMNS[i][0]
            if name.startswith(("วันที่", "รอบปีบัญชี")):
                cell.number_format = "dd/mm/yyyy"
            elif name.startswith("เวลา"):
                cell.number_format = "hh:mm"
    wd.freeze_panes = "A3"
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True)
    wd.add_data_validation(dv)
    dv.add("J3:J500")

    wb.save(path)


@dataclass
class DetailRow:
    business: str
    address: str
    phone: str
    fiscal_year: object
    signed_date: object
    date_from: object
    date_to: object
    time_from: object
    time_to: object
    hours: float
    work: str


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _to_date(v):
    """รับ date/datetime หรือข้อความ dd/mm/yyyy (ค.ศ. หรือ พ.ศ.)"""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for sep in ("/", "-", "."):
        parts = s.split(sep)
        if len(parts) == 3 and all(p.isdigit() for p in parts):
            d, m, y = (int(p) for p in parts)
            if y > 2400:
                y -= 543
            return dt.date(y, m, d)
    raise ValueError(f"อ่านวันที่ไม่ออก: {s!r} (ใช้รูปแบบ วว/ดด/ปปปป)")


def _to_time(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.time):
        return v
    if isinstance(v, (int, float)):  # Excel เก็บเวลาเป็นเศษของวัน
        minutes = round(float(v) * 24 * 60)
        return dt.time(minutes // 60, minutes % 60)
    s = str(v).strip().replace(".", ":")
    h, m = s.split(":")[:2]
    return dt.time(int(h), int(m))


def read_workbook(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in (SHEET_INFO, SHEET_DETAIL):
        if sheet not in wb.sheetnames:
            raise ValueError(f"ไม่พบชีท \"{sheet}\" ในไฟล์ {path}")
    info = {}
    for row in wb[SHEET_INFO].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            info[str(row[0]).strip()] = row[1] if len(row) > 1 else None

    lunch = info.get("หักพักกลางวัน (ชม./วัน)")
    lunch = 0.0 if lunch in (None, "") else float(lunch)

    rows = []
    ws = wb[SHEET_DETAIL]
    header = [_text(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    missing = [name for name, _, _ in DETAIL_COLUMNS if name not in header]
    if missing:
        raise ValueError(f"ชีท {SHEET_DETAIL} ไม่มีคอลัมน์: {', '.join(missing)}")
    col = {name: header.index(name) for name, _, _ in DETAIL_COLUMNS}
    for n, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        get = lambda name: r[col[name]] if col[name] < len(r) else None  # noqa: E731
        if not any(v not in (None, "") for v in r):
            continue
        t_from, t_to = _to_time(get("เวลาเริ่มต้น")), _to_time(get("เวลาสิ้นสุด"))
        hours = get("จำนวนชั่วโมง")
        if hours in (None, ""):
            if not (t_from and t_to):
                raise ValueError(f"ชีท {SHEET_DETAIL} แถว {n}: ไม่มีจำนวนชั่วโมง และไม่มีเวลาเริ่มต้น/สิ้นสุดให้คำนวณ")
            hours = (dt.datetime.combine(dt.date.min, t_to) - dt.datetime.combine(dt.date.min, t_from)).seconds / 3600
            if t_from <= dt.time(12) and t_to >= dt.time(13):
                hours -= lunch
        date_from = _to_date(get("วันที่เริ่ม"))
        if date_from is None:
            raise ValueError(f"ชีท {SHEET_DETAIL} แถว {n}: ไม่มีวันที่เริ่ม")
        rows.append(DetailRow(
            business=_text(get("ชื่อธุรกิจ")),
            address=_text(get("สถานที่ตั้ง")),
            phone=_text(get("เบอร์โทรศัพท์")),
            fiscal_year=get("รอบปีบัญชี"),
            signed_date=get("วันที่ผู้สอบบัญชีลงนามรับรองงบ"),
            date_from=date_from,
            date_to=_to_date(get("วันที่สิ้นสุด")),
            time_from=t_from,
            time_to=t_to,
            hours=float(hours),
            work=_text(get("ลักษณะการปฏิบัติงาน")),
        ))
    if not rows:
        raise ValueError(f"ชีท {SHEET_DETAIL} ไม่มีข้อมูล")
    return info, rows


# ---------------------------------------------------------------- formatting


def thai_date(d) -> str:
    """15 ม.ค. 2567"""
    if d in (None, ""):
        return ""
    if not isinstance(d, dt.date):
        try:
            d = _to_date(d)
        except ValueError:
            return _text(d)
    return f"{d.day}\u00a0{THAI_MONTHS_SHORT[d.month - 1]}\u00a0{d.year + 543}"  # ไม่ให้วันที่ขาดคนละบรรทัด


def thai_date_long(d) -> str:
    return f"{d.day} {THAI_MONTHS[d.month - 1]} พ.ศ. {d.year + 543}"


def fmt_hours(h: float) -> str:
    if float(h).is_integer():
        return f"{int(h):,}"
    return f"{h:,.2f}".rstrip("0").rstrip(".")


def fmt_time(t) -> str:
    return t.strftime("%H.%M") if t else ""


def span_ymd(start: dt.date, end: dt.date):
    """ระยะเวลาตั้งแต่ start ถึง end (นับวันสุดท้ายด้วย) เป็น ปี เดือน วัน"""
    end = end + dt.timedelta(days=1)
    y, m, d = end.year - start.year, end.month - start.month, end.day - start.day
    if d < 0:
        m -= 1
        prev_month_end = end.replace(day=1) - dt.timedelta(days=1)
        d += prev_month_end.day
    if m < 0:
        y -= 1
        m += 12
    return y, m, d


# ---------------------------------------------------------------- PDF helpers


@dataclass
class Run:
    """เส้นจุดไข่ปลาในฟอร์ม"""
    rect: pymupdf.Rect
    baseline: float


class FormPage:
    """หาตำแหน่งช่องจุดไข่ปลาจากข้อความในฟอร์ม แล้วพิมพ์ข้อความทับ"""

    def __init__(self, page: pymupdf.Page):
        self.page = page
        self.items = []
        # เรียงทีละ span ตามตำแหน่ง แต่ไม่เรียงตัวอักษรภายใน span
        # (สระบน/วรรณยุกต์ไทยมีพิกัด x ก่อนพยัญชนะ ถ้าเรียงตาม x ข้อความจะเพี้ยน)
        spans = []
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    cs = [(ch["c"], pymupdf.Rect(ch["bbox"]), span["font"]) for ch in span["chars"]]
                    if "".join(c[0] for c in cs).strip():
                        spans.append((span["origin"][1], span["bbox"][0], cs))
        spans.sort(key=lambda s: (round(s[0]), s[1]))
        self.lines = []  # [(baseline, [(char, rect, font)])]
        for base, _, cs in spans:
            if self.lines and abs(self.lines[-1][0] - base) < 3:
                self.lines[-1][1].extend(cs)
            else:
                self.lines.append((base, list(cs)))

    def _runs(self, line_no: int, start: int = 0):
        base, cs = self.lines[line_no]
        runs, i = [], start
        while i < len(cs):
            if cs[i][0] in ".…":
                j = i
                while j < len(cs) and cs[j][0] in ".…":
                    j += 1
                if j - i >= 3:
                    r = pymupdf.Rect(cs[i][1].x0, cs[i][1].y0, cs[j - 1][1].x1, cs[j - 1][1].y1)
                    runs.append(Run(r, base))
                i = j
            else:
                i += 1
        return runs

    def field(self, label: str, occurrence: int = 0, next_line: bool = False) -> Run:
        """ช่องจุดไข่ปลาที่อยู่ถัดจากข้อความ label (นับ occurrence ตามลำดับการอ่าน)"""
        seen = 0
        for n, (_, cs) in enumerate(self.lines):
            text = "".join(c[0] for c in cs)
            pos = text.find(label)
            while pos != -1:
                if seen == occurrence:
                    same = self._runs(n, pos + len(label))
                    if not next_line:
                        if not same:
                            break
                        return same[0]
                    below = next(runs for k in range(n + 1, len(self.lines)) if (runs := self._runs(k)))
                    if same:  # เช่น (ลงชื่อ)....  →  (.....) บรรทัดถัดไป ที่อยู่ตรงกัน
                        cx = (same[0].rect.x0 + same[0].rect.x1) / 2
                        return min(below, key=lambda r: abs((r.rect.x0 + r.rect.x1) / 2 - cx))
                    return below[0]
                seen += 1
                pos = text.find(label, pos + 1)
        raise LookupError(f"ไม่พบช่อง {label!r} ลำดับที่ {occurrence} ในหน้า {self.page.number + 1}")

    def boxes(self):
        """กล่องเลขประจำตัวฝึกหัดงาน (สัญลักษณ์ Wingdings)"""
        for _, cs in self.lines:
            found = [r for c, r, font in cs if font.startswith("Wingdings")]
            if found:
                return found
        raise LookupError("ไม่พบช่องเลขประจำตัวฝึกหัดงาน")

    def html(self, rect, text, size=FONT_SIZE, align="left"):
        """จองข้อความไว้พิมพ์ (บรรทัดละ 1 แถว ไม่ตัดบรรทัดเพิ่ม) วาดจริงตอน flush()"""
        self.items.append((pymupdf.Rect(rect), text, size, align))

    def write(self, run: Run, text: str, align: str = "center", size: float = FONT_SIZE):
        text = _text(text)
        if not text:
            return
        # ข้อความยาวเกินช่องจะถูกย่อตัวอักษรลง (ไม่ต่ำกว่าครึ่งหนึ่ง)
        room = run.rect.width - 2
        fit = min(size, size * room / text_width(text, size))
        if fit < size * 0.5:
            raise ValueError(f"ข้อความยาวเกินช่อง: {text!r}")
        # ให้ฐานตัวอักษรลอยเหนือเส้นจุดเล็กน้อย
        top = run.baseline - _ascent(fit, text) - 3
        self.html(pymupdf.Rect(run.rect.x0 + 1, top, run.rect.x1 - 1, top + fit * 4), text, size=fit, align=align)

    def flush(self):
        """วาดข้อความทั้งหน้าลงหน้ากระดาษเดียวแล้วซ้อนทับฟอร์ม (ฝังฟอนต์ครั้งเดียวต่อหน้า)"""
        if not self.items:
            return
        buf = io.BytesIO()
        writer = pymupdf.DocumentWriter(buf)
        device = writer.begin_page(self.page.rect)
        for rect, text, size, align in self.items:
            story = _story(text, size, align)
            more, _ = story.place(rect)
            if more:
                raise ValueError(f"ข้อความไม่พอดีช่อง: {text!r}")
            story.draw(device)
        writer.end_page()
        writer.close()
        overlay = pymupdf.open("pdf", buf.getvalue())
        self.page.show_pdf_page(self.page.rect, overlay, 0)
        self.items = []


def _story(text: str, size: float, align: str) -> pymupdf.Story:
    css = (
        "body {margin: 1px;}"
        f"* {{font-family: sans-serif; font-size: {size}pt; color: {INK}; margin: 0; padding: 0;"
        f" line-height: {LINE_HEIGHT};}}"
        f"div {{text-align: {align}; white-space: nowrap;}}"  # ตัดบรรทัดเองด้วย wrap_lines
    )
    body = "".join(f"<div>{_escape(line)}</div>" for line in text.split("\n"))
    return pymupdf.Story(html=body, user_css=css)


def _escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


_ASCENT = {}


def _is_thai(c: str) -> bool:
    return "\u0e00" <= c <= "\u0e7f"


def _ascent(size: float, text: str) -> float:
    """ระยะจากขอบบนกล่องข้อความถึงฐานบรรทัด วัดจริงจาก MuPDF (ไทย/อังกฤษใช้คนละฟอนต์ ฐานบรรทัดต่างกัน)"""
    kind = (any(_is_thai(c) for c in text), any(not _is_thai(c) and not c.isspace() for c in text))
    sample = {(True, False): "ก", (False, True): "0"}.get(kind, "ก0")
    if (size, sample) not in _ASCENT:
        doc = pymupdf.open()
        page = doc.new_page()
        fp = FormPage(page)
        fp.html(pymupdf.Rect(0, 0, 200, 100), sample, size=size)
        fp.flush()
        spans = [sp for b in page.get_text("rawdict")["blocks"] for ln in b["lines"] for sp in ln["spans"]]
        _ASCENT[(size, sample)] = max(sp["origin"][1] for sp in spans)
    return _ASCENT[(size, sample)]


_FONTS = {}
THAI_MARKS = "ัำิีึืฺุู็่้๊๋์ํ๎"
LEADING_VOWELS = "เแโใไ"

try:  # ตัดคำภาษาไทยด้วยพจนานุกรม (ถ้าติดตั้ง pythainlp ไว้)
    from pythainlp.tokenize import word_tokenize as _thai_tokenize
except ImportError:
    _thai_tokenize = None


def text_width(s: str, size: float) -> float:
    if not _FONTS:
        _FONTS["thai"] = pymupdf.Font(script=19)  # UCDN_SCRIPT_THAI → Noto Serif Thai ที่ฝังมากับ MuPDF
        _FONTS["latin"] = pymupdf.Font("helv")  # Nimbus Sans มีความกว้างเท่า Helvetica
    return sum(_FONTS["thai" if _is_thai(c) else "latin"].text_length(c, size) for c in s)


def _tokens(text: str):
    """แบ่งข้อความเป็นชิ้นที่ตัดบรรทัดได้ ภาษาไทยไม่เว้นวรรคระหว่างคำจึงต้องตัดคำก่อน"""
    out = []
    for chunk in text.split("\u00a0"):  # ช่องว่างแบบไม่ตัดบรรทัด: ต่อชิ้นหน้า-หลังเข้าด้วยกัน
        if _thai_tokenize is not None:
            toks = _thai_tokenize(chunk, keep_whitespace=True)
        else:  # ไม่มี pythainlp: ตัดที่ช่องว่าง และหน้าสระ เ แ โ ใ ไ
            toks = []
            for ch in chunk:
                if not toks or ch == " " or toks[-1] == " " or ch in LEADING_VOWELS:
                    toks.append(ch)
                else:
                    toks[-1] += ch
        if out and toks:
            out[-1] += "\u00a0" + toks.pop(0)
        elif out:
            out[-1] += "\u00a0"
        out.extend(toks)
    return out


def _clusters(word: str):
    """ตัวอักษรพร้อมสระบน/ล่างและวรรณยุกต์ (หน่วยเล็กสุดที่ตัดได้โดยไม่ทำให้สระลอย)"""
    out = []
    for ch in word:
        if out and ch in THAI_MARKS:
            out[-1] += ch
        else:
            out.append(ch)
    return out


def wrap_lines(text: str, widths, size: float):
    """ตัดข้อความเป็นบรรทัด บรรทัดที่ i กว้างไม่เกิน widths[i] (บรรทัดเกินจำนวนใช้ความกว้างสุดท้าย)"""
    lines = []
    for para in _text(text).split("\n"):
        line = ""
        for tok in _tokens(para.strip()):
            width = widths[min(len(lines), len(widths) - 1)]
            if text_width(line + tok, size) <= width:
                line += tok
                continue
            if line.strip():
                lines.append(line.rstrip())
                width = widths[min(len(lines), len(widths) - 1)]
            line = "" if tok == " " else tok
            while text_width(line, size) > width:  # คำเดียวยาวเกินบรรทัด
                head = ""
                for cl in _clusters(line):
                    if text_width(head + cl, size) > width and head:
                        break
                    head += cl
                lines.append(head)
                line = line[len(head):]
                width = widths[min(len(lines), len(widths) - 1)]
        if line.strip() or not lines:
            lines.append(line.rstrip())
    return lines


# ---------------------------------------------------------------- filling


@dataclass
class Data:
    trainee_id: str
    trainee: str
    auditor: str
    auditor_reg: str
    office: dict
    period_from: dt.date
    period_to: dt.date
    ps2_place: str
    ps2_date: dt.date
    ps3_place: str
    ps3_date: dt.date
    capacity: str
    opinion: str
    rows: list
    total_hours: float


def build_data(info: dict, rows: list) -> Data:
    def need(key):
        v = _text(info.get(key))
        if not v:
            raise ValueError(f"ชีท {SHEET_INFO}: ยังไม่ได้กรอก \"{key}\"")
        return v

    trainee_id = need("เลขประจำตัวฝึกหัดงาน").replace("-", "").replace(" ", "")
    if not trainee_id.isdigit() or len(trainee_id) > 7:
        raise ValueError(f"เลขประจำตัวฝึกหัดงานต้องเป็นตัวเลขไม่เกิน 7 หลัก: {trainee_id!r}")

    rows = sorted(rows, key=lambda r: (r.date_from, r.time_from or dt.time.min))
    period_from = _to_date(info.get("ฝึกตั้งแต่วันที่")) or rows[0].date_from
    period_to = _to_date(info.get("ฝึกถึงวันที่")) or max(r.date_to or r.date_from for r in rows)
    ps2_date = _to_date(info.get("ผส.2 วันที่")) or dt.date.today()

    office_keys = ["สำนักงาน", "เลขที่", "หมู่ที่", "ซอย", "ถนน", "ตำบล/แขวง", "อำเภอ/เขต",
                   "จังหวัด", "รหัสไปรษณีย์", "โทรศัพท์", "โทรสาร", "E-mail"]
    return Data(
        trainee_id=trainee_id.zfill(7),
        trainee=need("ชื่อผู้ฝึกหัดงาน"),
        auditor=need("ชื่อผู้สอบบัญชี"),
        auditor_reg=need("เลขทะเบียนผู้สอบบัญชี"),
        office={k: _text(info.get(k)) for k in office_keys},
        period_from=period_from,
        period_to=period_to,
        ps2_place=_text(info.get("ผส.2 เขียนที่")),
        ps2_date=ps2_date,
        ps3_place=_text(info.get("ผส.3 เขียนที่")) or _text(info.get("ผส.2 เขียนที่")),
        ps3_date=_to_date(info.get("ผส.3 วันที่")) or ps2_date,
        capacity=_text(info.get("ฝึกในฐานะ")),
        opinion=_text(info.get("ความเห็นของผู้ให้การฝึก")),
        rows=rows,
        total_hours=sum(r.hours for r in rows),
    )


def fill_id_boxes(fp: FormPage, trainee_id: str):
    for digit, box in zip(trainee_id, fp.boxes()):
        # กล่องที่มองเห็นแคบกว่า bbox ของตัวอักษร Wingdings
        rect = pymupdf.Rect(box.x0 + 2, box.y0 + 7.5, box.x0 + 19, box.y1 + 10)
        fp.html(rect, digit, size=13, align="center")


def write_date(fp: FormPage, d: dt.date, occurrence: int):
    fp.write(fp.field("วันที่", occurrence), str(d.day))
    fp.write(fp.field("เดือน", 0), THAI_MONTHS[d.month - 1])
    fp.write(fp.field("พ.ศ.", 0), str(d.year + 543))


def fill_ps2(fp: FormPage, data: Data):
    o = data.office
    fill_id_boxes(fp, data.trainee_id)
    fp.write(fp.field("เขียนที่"), data.ps2_place)
    write_date(fp, data.ps2_date, occurrence=1)  # วันที่ ลำดับ 0 อยู่ในกรอบเจ้าหน้าที่
    fp.write(fp.field("นาย/นางสาว/นาง/ยศ"), data.trainee)
    fp.write(fp.field("ผู้สอบบัญชีรับอนุญาต ชื่อ"), data.auditor)
    fp.write(fp.field("เลขทะเบียน", 0), data.auditor_reg)
    fp.write(fp.field("สำนักงาน"), o["สำนักงาน"])
    fp.write(fp.field("ตั้งอยู่เลขที่"), o["เลขที่"])
    fp.write(fp.field("หมู่ที่"), o["หมู่ที่"])
    fp.write(fp.field("ตรอก/ซอย"), o["ซอย"])
    fp.write(fp.field("ถนน"), o["ถนน"])
    fp.write(fp.field("ตำบล/แขวง"), o["ตำบล/แขวง"])
    fp.write(fp.field("อำเภอ/เขต"), o["อำเภอ/เขต"])
    fp.write(fp.field("จังหวัด"), o["จังหวัด"])
    fp.write(fp.field("รหัสไปรษณีย์"), o["รหัสไปรษณีย์"])
    fp.write(fp.field("โทรศัพท์"), o["โทรศัพท์"])
    fp.write(fp.field("โทรสาร"), o["โทรสาร"])
    fp.write(fp.field("ตั้งแต่วันที่"), thai_date_long(data.period_from))
    fp.write(fp.field("ถึงวันที่", next_line=True), thai_date_long(data.period_to))
    fp.write(fp.field("ทั้งสิ้น จำนวน"), fmt_hours(data.total_hours))
    fp.write(fp.field("(ลงชื่อ)", 0, next_line=True), data.trainee)
    fp.write(fp.field("3. ข้าพเจ้า"), data.auditor)
    fp.write(fp.field("เลขทะเบียน", 1), data.auditor_reg)
    fp.write(fp.field("รายงานการฝึกหัดงานสอบบัญชีของ"), data.trainee)
    fp.write(fp.field("(ลงชื่อ)", 1, next_line=True), data.auditor)
    fp.write(fp.field("เลขทะเบียน", 2), data.auditor_reg)


def fill_ps3(fp: FormPage, data: Data):
    o = data.office
    fill_id_boxes(fp, data.trainee_id)
    fp.write(fp.field("เขียนที่"), data.ps3_place)
    write_date(fp, data.ps3_date, occurrence=1)
    fp.write(fp.field("1. ข้าพเจ้า นาย/นาง/นางสาว/ยศ"), data.auditor)
    fp.write(fp.field("เลขทะเบียน"), data.auditor_reg)
    fp.write(fp.field("สำนักงานชื่อ"), o["สำนักงาน"])
    fp.write(fp.field("ตั้งอยู่เลขที่"), o["เลขที่"])
    fp.write(fp.field("หมู่ที่"), o["หมู่ที่"])
    fp.write(fp.field("ซอย"), o["ซอย"])
    fp.write(fp.field("ถนน"), o["ถนน"])
    fp.write(fp.field("ตำบล/แขวง"), o["ตำบล/แขวง"])
    fp.write(fp.field("อำเภอ/เขต"), o["อำเภอ/เขต"])
    fp.write(fp.field("จังหวัด"), o["จังหวัด"])
    fp.write(fp.field("รหัสไปรษณีย์"), o["รหัสไปรษณีย์"])
    fp.write(fp.field("โทรศัพท์"), o["โทรศัพท์"])
    fp.write(fp.field("โทรสาร"), o["โทรสาร"])
    fp.write(fp.field("E-mail"), o["E-mail"], size=FONT_SIZE - 1)
    fp.write(fp.field("ว่า นาย/นาง/นางสาว/ยศ", 0), data.trainee)
    fp.write(fp.field("ในฐานะ"), data.capacity)
    fp.write(fp.field("ตั้งแต่วันที่"), thai_date(data.period_from))
    fp.write(fp.field("ถึงวันที่"), thai_date(data.period_to))
    y, m, d = span_ymd(data.period_from, data.period_to)
    fp.write(fp.field("รวมเป็นเวลา"), str(y))
    fp.write(fp.field("ปี"), str(m))
    fp.write(fp.field("เดือน", 1), str(d))
    fp.write(fp.field("ว่า นาย/นาง/นางสาว/ยศ", 1), data.trainee)
    if data.opinion:
        first = fp.field("การสอบบัญชีและ")
        second = fp.field("การสอบบัญชีและ", next_line=True)
        parts = wrap_lines(data.opinion, [first.rect.width - 4, second.rect.width - 4], FONT_SIZE)
        if len(parts) > 2:  # ยาวเกินสองบรรทัด: บรรทัดสุดท้ายรับส่วนที่เหลือทั้งหมด (จะถูกย่อตัวอักษรลง)
            parts = parts[:1] + [" ".join(parts[1:])]
        fp.write(first, parts[0], align="left")
        if len(parts) > 1:
            fp.write(second, parts[1], align="left")
    fp.write(fp.field("(ลงชื่อ)", 0, next_line=True), data.auditor)


# ตารางรายละเอียด: ขอบคอลัมน์ (x) และช่วงแนวตั้งของตัวตาราง วัดจากฟอร์มเปล่า
TABLES = {
    PAGE_DETAIL_SHORT: dict(
        cols=[24.1, 73.8, 272.3, 385.7, 470.7, 571.2],
        body=(126.7, 629.8),
        total=pymupdf.Rect(385.7, 630.3, 470.3, 651.3),
    ),
    PAGE_DETAIL_LONG: dict(
        cols=[21.5, 64.0, 247.7, 325.4, 373.9, 423.4, 475.3, 573.9],
        body=(153.3, 632.9),
        total=pymupdf.Rect(423.4, 633.3, 474.8, 653.3),
    ),
}


def business_cell(r: DetailRow, long_form: bool) -> str:
    lines = [r.business]
    if r.address:
        lines.append(r.address)
    extra = []
    if r.phone:
        extra.append(f"โทร. {r.phone}")
    if r.fiscal_year:
        extra.append(f"รอบปีบัญชี {thai_date(r.fiscal_year)}")
    if extra:
        lines.append(" / ".join(extra))
    if long_form and r.signed_date:
        lines.append(f"ลงนามรับรองงบ {thai_date(r.signed_date)}")
    return "\n".join(lines)


def table_cells(r: DetailRow, seq: int, long_form: bool, show_business: bool):
    biz = business_cell(r, long_form) if show_business else ""
    if long_form:
        return [str(seq), biz, thai_date(r.date_from), fmt_time(r.time_from), fmt_time(r.time_to),
                fmt_hours(r.hours), r.work]
    period = thai_date(r.date_from)
    if r.date_to and r.date_to != r.date_from:
        period += f" -\n{thai_date(r.date_to)}"
    return [str(seq), biz, period, fmt_hours(r.hours), r.work]


def wrap_cell(text: str, width: float):
    """ตัดบรรทัดข้อความในช่องตาราง คืน (ข้อความที่ตัดแล้ว, ความสูงที่ต้องใช้)"""
    if not text:
        return "", 0.0
    lines = wrap_lines(text, [width], TABLE_FONT_SIZE)
    return "\n".join(lines), len(lines) * TABLE_FONT_SIZE * LINE_HEIGHT


def paginate_rows(rows, long_form: bool):
    """แบ่งแถวลงหน้า (หน้าละเท่าที่ตารางรับได้) คืน [[(cells, height), ...], ...]"""
    t = TABLES[PAGE_DETAIL_LONG if long_form else PAGE_DETAIL_SHORT]
    widths = [b - a - 4 for a, b in zip(t["cols"], t["cols"][1:])]
    room = t["body"][1] - t["body"][0] - 4
    pages, current, used, last_biz = [], [], 0.0, None
    for seq, r in enumerate(rows, start=1):
        for attempt in range(2):
            biz_key = (r.business, r.address)
            cells = table_cells(r, seq, long_form, show_business=biz_key != last_biz or not current)
            cells, heights = zip(*(wrap_cell(c, w) for c, w in zip(cells, widths)))
            h = max(heights) + 4
            if used + h <= room or not current:
                break
            pages.append(current)  # ขึ้นหน้าใหม่ แล้วคำนวณใหม่ (หน้าใหม่ต้องแสดงชื่อธุรกิจซ้ำ)
            current, used = [], 0.0
        current.append((cells, h))
        used += h
        last_biz = biz_key
    if current:
        pages.append(current)
    return pages


def fill_detail(fp: FormPage, data: Data, page_rows, long_form: bool, subtotal: float, is_last: bool):
    t = TABLES[PAGE_DETAIL_LONG if long_form else PAGE_DETAIL_SHORT]
    fp.write(fp.field("ยศ"), data.trainee)
    cols = t["cols"]
    y = t["body"][0] + 2
    aligns = (["center", "left", "center", "center", "center", "center", "left"] if long_form
              else ["center", "left", "center", "center", "left"])
    for n, (cells, h) in enumerate(page_rows):
        for text, x0, x1, align in zip(cells, cols, cols[1:], aligns):
            if text:
                fp.html(pymupdf.Rect(x0 + 2, y, x1 - 2, y + h + 20), text, size=TABLE_FONT_SIZE, align=align)
        y += h
        if n < len(page_rows) - 1:
            fp.page.draw_line((cols[0] + 0.5, y - 1), (cols[-1] - 0.5, y - 1), color=(0.75, 0.75, 0.75), width=0.4)
    total = t["total"]
    label = fmt_hours(subtotal) if is_last else f"{fmt_hours(subtotal)} (ยกไป)"
    fp.html(pymupdf.Rect(total.x0 + 1, total.y0 + 3, total.x1 - 1, total.y1 + 20), label,
            size=TABLE_FONT_SIZE if is_last else TABLE_FONT_SIZE - 2, align="center")
    fp.write(fp.field("(ลงชื่อ)", 0, next_line=True), data.trainee, size=FONT_SIZE - 1)
    fp.write(fp.field("(ลงชื่อ)", 1, next_line=True), data.auditor, size=FONT_SIZE - 1)


def generate(data: Data, out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    blank = pymupdf.open(BLANK_FORM)
    long_form = data.total_hours > HOUR_LIMIT
    detail_page = PAGE_DETAIL_LONG if long_form else PAGE_DETAIL_SHORT
    pages = paginate_rows(data.rows, long_form)
    safe_name = "".join(c for c in data.trainee if c not in '\\/:*?"<>|').strip()

    ps2 = pymupdf.open()
    ps2.insert_pdf(blank, from_page=PAGE_PS2, to_page=PAGE_PS2)
    fp = FormPage(ps2[0])
    fill_ps2(fp, data)
    fp.flush()
    running = 0.0
    for i, page_rows in enumerate(pages):
        ps2.insert_pdf(blank, from_page=detail_page, to_page=detail_page)
        seqs = [int(cells[0]) for cells, _ in page_rows]
        running += sum(data.rows[s - 1].hours for s in seqs)
        fp = FormPage(ps2[-1])
        fill_detail(fp, data, page_rows, long_form, running, i == len(pages) - 1)
        fp.flush()
    ps2_path = out_dir / f"แบบ ผส.2 - {safe_name}.pdf"
    ps2.save(ps2_path, garbage=3, deflate=True)

    ps3 = pymupdf.open()
    ps3.insert_pdf(blank, from_page=PAGE_PS3, to_page=PAGE_PS3)
    fp = FormPage(ps3[0])
    fill_ps3(fp, data)
    fp.flush()
    ps3_path = out_dir / f"แบบ ผส.3 - {safe_name}.pdf"
    ps3.save(ps3_path, garbage=3, deflate=True)
    return ps2_path, ps3_path, long_form, len(pages)


def main(argv=None):
    ap = argparse.ArgumentParser(description="เติมแบบ ผส.2 / ผส.3 สภาวิชาชีพบัญชี จากไฟล์ Excel")
    ap.add_argument("excel", type=Path, help="ไฟล์ข้อมูล .xlsx")
    ap.add_argument("-o", "--output", type=Path, default=Path("output"), help="โฟลเดอร์เก็บ PDF (ค่าเริ่มต้น: output)")
    ap.add_argument("--make-template", action="store_true", help="สร้างไฟล์ Excel ตัวอย่างไว้กรอกข้อมูล")
    args = ap.parse_args(argv)

    if args.make_template:
        if args.excel.exists():
            sys.exit(f"มีไฟล์ {args.excel} อยู่แล้ว ไม่เขียนทับ")
        make_template(args.excel)
        print(f"สร้างไฟล์ {args.excel} แล้ว กรอกข้อมูลแล้วรัน: python fill_forms.py {args.excel}")
        return

    try:
        info, rows = read_workbook(args.excel)
        data = build_data(info, rows)
        ps2_path, ps3_path, long_form, n_pages = generate(data, args.output)
    except FileNotFoundError:
        sys.exit(f"ไม่พบไฟล์ {args.excel} (สร้างไฟล์ใหม่ด้วย --make-template)")
    except (ValueError, LookupError) as e:
        sys.exit(f"ผิดพลาด: {e}")

    kind = "มากกว่า" if long_form else "ไม่เกิน"
    print(f"ชั่วโมงฝึกหัดงานรวม {fmt_hours(data.total_hours)} ชม. ({kind} 1,500) "
          f"ใช้ตารางรายละเอียด {n_pages} หน้า")
    print(f"  {ps2_path}")
    print(f"  {ps3_path}")


if __name__ == "__main__":
    main()
